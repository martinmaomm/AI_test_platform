#!/usr/bin/env python3
"""Offline browser integration test for the project knowledge workspace.

It serves the already-built Vue distribution and Django on one random loopback
origin.  All persistence is a temporary SQLite file; external sockets are
blocked in Python and all browser requests outside that origin are aborted.
"""
from __future__ import annotations

import json
import argparse
from concurrent.futures import ThreadPoolExecutor
import mimetypes
import os
from pathlib import Path
import re
import socket
import sys
import tempfile
import threading
import traceback
from urllib.parse import urlsplit
import zipfile
from types import SimpleNamespace
from unittest.mock import patch
from wsgiref.simple_server import WSGIRequestHandler, make_server


BACKEND_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIST = BACKEND_DIR.parent / "frontend" / "dist"
CHROME = Path(os.environ.get("AITS_TEST_CHROME_EXECUTABLE", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"))
UUID_RE = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", re.I)
OFFLINE_JWT_SIGNING_KEY = "project-knowledge-browser-offline-only-signing-key"


class _NoMigrations(dict):
    """Make migrate --run-syncdb create a disposable schema without migration I/O."""

    def __contains__(self, key):
        # MigrationLoader checks membership before it calls ``.get``.  Claim
        # every app so no installed app falls back to its on-disk migrations.
        return True

    def get(self, key, default=None):  # Django calls ``.get`` for every app label.
        return None

    def __getitem__(self, key):
        return None


class _QuietHandler(WSGIRequestHandler):
    def log_message(self, format, *args):
        return


def _loopback_only(original):
    def guarded(sock, address):
        host = address[0] if isinstance(address, tuple) else ""
        if host in {"127.0.0.1", "::1", "localhost"}:
            return original(sock, address)
        raise RuntimeError(f"离线知识库浏览器测试禁止外部 socket 连接: {address!r}")
    return guarded


def _static_or_django(django_app):
    def application(environ, start_response):
        path = environ.get("PATH_INFO", "/")
        if path.startswith("/api/") or path.startswith("/admin/") or path.startswith("/media/"):
            return django_app(environ, start_response)
        relative = path.lstrip("/")
        candidate = (FRONTEND_DIST / relative).resolve()
        try:
            candidate.relative_to(FRONTEND_DIST.resolve())
        except ValueError:
            start_response("403 Forbidden", [("Content-Type", "text/plain; charset=utf-8")])
            return [b"forbidden"]
        if not candidate.is_file():
            candidate = FRONTEND_DIST / "index.html"
        content_type = mimetypes.guess_type(str(candidate))[0] or "application/octet-stream"
        start_response("200 OK", [("Content-Type", content_type), ("Content-Length", str(candidate.stat().st_size))])
        return [candidate.read_bytes()]
    return application


def _first_chunk_id(messages) -> str:
    joined = json.dumps(messages, ensure_ascii=False)
    match = UUID_RE.search(joined)
    if not match:
        raise AssertionError("假 LLM 未能从真实工作流 prompt 找到 chunk UUID")
    return match.group(0)


def _fake_stream_call(model_config_id, messages, on_chunk, check_active, remaining_seconds):
    check_active()
    system = next((str(item.get("content", "")) for item in messages if isinstance(item, dict) and item.get("role") == "system"), "")
    chunk_id = _first_chunk_id(messages)
    if '"test_points"' in system:
        response = {
            "test_points": [{
                "title": "停用角色登录校验", "kind": "权限", "description": "校验停用角色无法登录",
                "source_ids": [chunk_id], "pending_questions": [],
            }],
        }
    elif '"cases"' in system:
        response = {
            "cases": [{
                "title": "停用角色不可登录", "module": "账号", "test_type": "权限",
                "preconditions": ["存在已停用角色"], "test_data": "停用账号",
                "steps": [{"action": "使用停用账号登录", "expected": "系统拒绝登录"}],
                "sources": [chunk_id], "test_point_ids": [f"tp-{chunk_id}-1"], "pending_questions": [],
            }],
        }
    elif '"answer"' in system:
        response = {
            "answer": "依据当前资料，角色停用后不可登录。",
            "result_type": "supported", "source_ids": [chunk_id], "warnings": [],
        }
    else:
        raise AssertionError(f"假 LLM 收到未知知识库 prompt: {system[:100]}")
    rendered = json.dumps(response, ensure_ascii=False)
    if on_chunk:
        on_chunk(rendered)
    return rendered


def _fake_search(project_id, revision_ids, query, top_k=8):
    from project_knowledge.models import KnowledgeChunk

    rows = KnowledgeChunk.objects.select_related("revision__document").filter(
        revision__document__project_id=project_id,
        revision_id__in=[str(value) for value in revision_ids],
        revision__document__is_active=True,
        revision__document__is_deleted=False,
        revision__index_status="ready",
    ).order_by("ordinal")[:top_k]
    return [{
        "id": str(row.id), "content": row.content, "file_name": row.revision.uploaded_file.original_name,
        "document_id": row.revision.document_id, "revision_id": str(row.revision_id),
        "heading": row.heading, "location": row.location,
    } for row in rows]


def _inline_delay(task_id):
    from project_knowledge import runtime

    runtime.execute_task(task_id)
    return SimpleNamespace(id=f"inline-{task_id}")


def _bootstrap_database(root: Path):
    sys.path[:0] = [str(BACKEND_DIR), str(BACKEND_DIR / "apps")]
    os.environ["DJANGO_SETTINGS_MODULE"] = "aits_backend.settings"
    os.environ["ANONYMIZED_TELEMETRY"] = "false"
    os.environ["MCP_USE_ANONYMIZED_TELEMETRY"] = "false"

    from aits_backend import settings as config

    config.DATABASES = {"default": {"ENGINE": "django.db.backends.sqlite3", "NAME": str(root / "knowledge.sqlite3")}}
    # Never derive synthetic browser JWTs from a secret loaded via the user's
    # environment.  This key exists only in this process and temporary SQLite.
    config.SECRET_KEY = OFFLINE_JWT_SIGNING_KEY
    if "SIGNING_KEY" in config.SIMPLE_JWT:
        config.SIMPLE_JWT = {**config.SIMPLE_JWT, "SIGNING_KEY": OFFLINE_JWT_SIGNING_KEY}
    config.MIGRATION_MODULES = _NoMigrations()
    config.CACHES = {"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}}
    config.CHANNEL_LAYERS = {"default": {"BACKEND": "channels.layers.InMemoryChannelLayer"}}
    config.CELERY_BROKER_URL = "memory://"
    config.CELERY_RESULT_BACKEND = "cache+memory://"
    config.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]
    config.MEDIA_ROOT = str(root / "media")
    config.STATIC_ROOT = str(root / "static")
    config.ALLOWED_HOSTS = ["127.0.0.1", "localhost"]
    config.PROJECT_KNOWLEDGE_ENABLED = True
    config.LOGGING = {"version": 1, "disable_existing_loggers": True}

    import django
    django.setup()
    from django.core.management import call_command
    call_command("migrate", run_syncdb=True, verbosity=0)

    from ai_core.models import LLMConfiguration, ModelType
    from django.contrib.auth import get_user_model
    from projects.models import Project

    user = get_user_model().objects.create_user(
        username="knowledge-browser", email="knowledge-browser@example.test", password="offline-only",
    )
    project = Project.objects.create(name="浏览器隔离知识库项目", project_type="web", created_by=user)
    model = LLMConfiguration.objects.create(
        model_type=ModelType.LLM, provider="openai", provider_name="fake", model_name="offline-json",
        is_active=True, api_key="", base_url="", created_by=user,
    )
    from rest_framework_simplejwt.tokens import AccessToken
    return user, project, model, str(AccessToken.for_user(user))


def _ui_flow(origin: str, root: Path, artifacts_dir: Path, token: str, user, project, model):
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError, sync_playwright

    console_errors, page_errors, api_failures = [], [], []
    api_requests, api_responses, option_payloads = [], [], []
    fixture = root / "requirements.md"
    fixture.write_text("# 登录规则\n角色停用后不可登录，系统必须拒绝登录。\n", encoding="utf-8")
    workspace_screenshot = artifacts_dir / "workspace.png"
    cases_screenshot = artifacts_dir / "manual-cases.png"
    qa_screenshot = artifacts_dir / "qa-sources-history.png"
    download_path = root / "manual-test-cases.xlsx"

    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(root / "chrome-profile"), executable_path=str(CHROME), headless=True,
            accept_downloads=True, viewport={"width": 1440, "height": 1100},
            args=["--no-first-run", "--no-default-browser-check"],
        )
        context.add_init_script(f"""
            localStorage.setItem('auth-store', JSON.stringify({{accessToken: {json.dumps(token)}, refreshToken: null,
              user: {{id: {user.id}, username: 'knowledge-browser', email: 'knowledge-browser@example.test'}}}}));
            localStorage.setItem('project-store', JSON.stringify({{currentProject: {{id: {project.id}, name: '浏览器隔离知识库项目', project_type: 'web'}}}}));
        """)

        def route(route, request):
            if request.url.startswith(origin + "/"):
                route.continue_()
            else:
                route.abort()

        context.route("**/*", route)
        page = context.new_page()
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        page.on("pageerror", lambda error: page_errors.append(str(error)))

        def trace_url(url):
            return urlsplit(url).path

        def observe_request(request):
            if request.url.startswith(origin + "/api/"):
                # Deliberately record only method and path: never headers,
                # request bodies, query strings, or the synthetic JWT.
                api_requests.append(f"{request.method} {trace_url(request.url)}")

        def observe_response(response):
            if response.url.startswith(origin + "/api/"):
                trace = f"{response.status} {trace_url(response.url)}"
                api_responses.append(trace)
                if response.status >= 400:
                    api_failures.append(trace)
            if "/knowledge/options/" in response.url:
                option_payloads.append(response.text())

        page.on("request", observe_request)
        page.on("response", observe_response)

        def assert_rendered_draft_text(text, label):
            locator = page.locator(".draft-item").get_by_text(text, exact=True)
            locator.wait_for(state="visible")
            font_size = locator.evaluate("element => getComputedStyle(element).fontSize")
            box = locator.bounding_box()
            if float(font_size.removesuffix("px")) <= 0 or not box or box["height"] <= 0:
                raise AssertionError(f"{label} 未实际渲染: font-size={font_size}, boundingBox={box}")

        def wait_for_toasts():
            try:
                page.wait_for_function(
                    """() => ![...document.querySelectorAll('.el-message')].some((element) => {
                        const style = getComputedStyle(element);
                        const box = element.getBoundingClientRect();
                        return style.display !== 'none' && style.visibility !== 'hidden' && box.width > 0 && box.height > 0;
                    })""",
                    timeout=5_000,
                )
            except PlaywrightTimeoutError as error:
                raise AssertionError("截图前 .el-message 未在 5 秒内消失") from error

        try:
            try:
                with page.expect_response("**/knowledge/options/", timeout=30_000) as option_response_info:
                    page.goto(f"{origin}/web-testing/knowledge-base", wait_until="domcontentloaded")
                option_payloads.append(option_response_info.value.text())
            except PlaywrightTimeoutError as error:
                raise AssertionError(
                    "等待知识库 options 响应超时；"
                    f"requests={api_requests}；responses={api_responses}；"
                    f"pageerrors={page_errors}；console={console_errors}"
                ) from error
            page.get_by_role("heading", name="项目知识库").wait_for()

            page.get_by_role("button", name="上传资料").click()
            dialog = page.get_by_role("dialog", name="上传项目资料")
            dialog.locator("input[type=file]").set_input_files(str(fixture))
            dialog.get_by_role("button", name="上传", exact=True).click()
            page.get_by_text("资料已上传，尚未处理。", exact=True).wait_for()
            page.get_by_role("button", name="处理文档", exact=True).click()
            page.get_by_text("可用于问答", exact=True).wait_for()

            page.get_by_role("button", name="章节 / 来源", exact=True).click()
            page.get_by_text("角色停用后不可登录", exact=False).wait_for()
            wait_for_toasts()
            page.screenshot(path=str(workspace_screenshot), animations="disabled")
            page.keyboard.press("Escape")

            page.get_by_role("tab", name="手工测试用例").click()
            page.get_by_role("button", name="生成手工用例", exact=True).click()
            dialog = page.get_by_role("dialog", name="生成手工测试用例")
            dialog.get_by_placeholder("例如：用户管理的新增、编辑及权限校验").fill("验证停用角色登录规则")
            dialog.locator(".el-form-item").filter(has_text="主需求资料").locator(".el-select").click()
            page.get_by_text("requirements.md", exact=True).last.click()
            dialog.locator(".form-grid .el-select").last.get_by_text("fake - offline-json", exact=True).wait_for()
            dialog.get_by_role("button", name="开始生成", exact=True).click()
            page.get_by_text("停用角色不可登录", exact=True).wait_for()
            assert_rendered_draft_text("使用停用账号登录", "草稿步骤")
            assert_rendered_draft_text("系统拒绝登录", "草稿预期")
            page.get_by_role("checkbox", name="停用角色不可登录").check()
            page.get_by_role("button", name=re.compile(r"保存选中")).click()
            page.get_by_text("已保存 1 条手工用例", exact=True).wait_for()

            page.get_by_role("button", name="编辑", exact=True).click()
            dialog = page.get_by_role("dialog", name="编辑手工测试用例")
            dialog.locator("input").first.fill("停用角色不可登录（已编辑）")
            dialog.get_by_role("button", name="保存", exact=True).click()
            page.get_by_text("手工用例已更新", exact=True).wait_for()
            page.locator(".el-table .el-checkbox").first.click()
            with page.expect_download() as download_info:
                page.get_by_role("button", name="导出选中 XLSX", exact=True).click()
            download_info.value.save_as(str(download_path))
            if not download_path.exists() or download_path.stat().st_size == 0:
                raise AssertionError("浏览器导出的 XLSX 文件为空")
            if not zipfile.is_zipfile(download_path):
                raise AssertionError("浏览器导出的文件不是有效 XLSX/ZIP")
            from openpyxl import load_workbook
            workbook = load_workbook(download_path, read_only=True, data_only=False)
            try:
                sheet = workbook["手工测试用例"]
                header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                exported = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
                if header[:7] != ("标题", "模块", "类型", "前置条件", "测试数据", "测试步骤", "预期结果"):
                    raise AssertionError(f"XLSX 表头不符合导出契约: {header}")
                if exported[0] != "停用角色不可登录（已编辑）" or "1. 使用停用账号登录" not in exported[5] or "1. 系统拒绝登录" not in exported[6]:
                    raise AssertionError(f"XLSX 用例数据或步骤/预期双列错误: {exported[:7]}")
            finally:
                workbook.close()
            page.locator(".draft-item").scroll_into_view_if_needed()
            wait_for_toasts()
            page.screenshot(path=str(cases_screenshot), animations="disabled")

            page.get_by_role("tab", name="知识问答").click()
            page.get_by_role("button", name="新建", exact=True).click()
            page.get_by_placeholder("针对当前项目资料提问…").fill("角色停用后能否登录？")
            page.locator(".question-box .el-select").get_by_text("fake - offline-json", exact=True).wait_for()
            page.get_by_role("button", name="提问", exact=True).click()
            page.get_by_text("依据当前资料，角色停用后不可登录。", exact=True).wait_for()
            _verify_conversation_history(page, artifacts_dir)
            page.get_by_role("button", name=re.compile(r"requirements\.md.*登录规则")).first.click()
            page.get_by_label("来源原文").get_by_text("角色停用后不可登录，系统必须拒绝登录。", exact=False).wait_for()
            wait_for_toasts()
            page.screenshot(path=str(qa_screenshot), animations="disabled")
            page.keyboard.press('Escape')
            _verify_conversation_deletion(page, artifacts_dir)
        except Exception:
            page.screenshot(path=str(artifacts_dir / "failure.png"), animations="disabled")
            raise
        finally:
            context.close()

    if console_errors:
        raise AssertionError(f"浏览器 console error: {console_errors}")
    if api_failures:
        raise AssertionError(f"浏览器 API 失败: {api_failures}")


def _verify_conversation_history(page, artifacts_dir):
    """New/switch/reload must preserve saved history and each page-local draft."""
    from playwright.sync_api import expect

    first_question = '角色停用后能否登录？'
    answer = '依据当前资料，角色停用后不可登录。'
    question_box = page.get_by_placeholder('针对当前项目资料提问…')
    original = page.locator('.conversation-panel .el-menu-item').filter(has_text=first_question)
    expect(original).to_be_visible()
    question_box.fill('原会话未发送的草稿')
    page.get_by_role('button', name='新建', exact=True).click()
    expect(page.locator('.conversation-panel .el-menu-item')).to_have_count(2)
    expect(question_box).to_have_value('')
    expect(page.locator('.chat-message')).to_have_count(0)
    question_box.fill('第二个会话的草稿')
    original.click()
    expect(question_box).to_have_value('原会话未发送的草稿')
    expect(page.get_by_text(answer, exact=True)).to_be_visible()

    # Delay the HTTP acknowledgement after the real temporary API has accepted
    # the question. No real worker, external model, or socket is involved.
    held = {}

    def hold_post(route):
        if route.request.method != 'POST':
            route.continue_()
            return
        held['response'] = route.fetch()
        held['route'] = route

    pattern = '**/knowledge/conversations/*/messages/'
    page.route(pattern, hold_post)
    with patch('project_knowledge.tasks.execute_knowledge_task.delay', return_value=SimpleNamespace(id='offline-delayed-answer')):
        question_box.fill('如何验证角色停用后的登录限制？')
        page.get_by_role('button', name='提问', exact=True).click()
        page.get_by_role('button', name='新建', exact=True).click()
        expect(page.locator('.conversation-panel .el-menu-item')).to_have_count(3)
        expect(question_box).to_have_value('')
        question_box.fill('第三个会话仍需保留的草稿')
        assert 'response' in held, 'Question was not accepted by the isolated API'
        task_id = held['response'].json()['data']['task']['id']
        held['route'].fulfill(response=held['response'])
        page.unroute(pattern, hold_post)
        expect(question_box).to_have_value('第三个会话仍需保留的草稿')
        expect(page.locator('.chat-message')).to_have_count(0)
        original.click()
        expect(page.locator('.chat-panel .task-card')).to_contain_text('排队中')
        expect(original.get_by_role('button', name=re.compile(r'^删除会话'))).to_be_disabled()
        expect(question_box).to_have_value('')
        # Sync Playwright keeps an event loop on this thread. Run the fake
        # worker's synchronous ORM calls outside that loop, as a real worker does.
        with ThreadPoolExecutor(max_workers=1) as executor:
            executor.submit(_inline_delay, task_id).result(timeout=15)
        expect(page.locator('.chat-panel .task-card')).to_contain_text('已完成', timeout=15000)
        expect(page.locator('.chat-message.assistant pre').filter(has_text=answer)).to_have_count(2)

    page.locator('.conversation-panel .el-menu-item').nth(1).click()
    expect(question_box).to_have_value('第二个会话的草稿')
    page.locator('.conversation-panel .el-menu-item').nth(0).click()
    expect(question_box).to_have_value('第三个会话仍需保留的草稿')
    page.reload()
    page.get_by_role('tab', name='知识问答').click()
    original.click()
    expect(page.locator('.chat-message')).to_have_count(4)
    expect(page.locator('.chat-message.assistant pre').filter(has_text=answer)).to_have_count(2)
    expect(page.locator('.chat-messages .el-loading-mask:visible')).to_have_count(0)
    page.screenshot(path=str(artifacts_dir / 'qa-restored-history.png'), animations='disabled')


def _verify_conversation_deletion(page, artifacts_dir):
    """Exercise only disposable conversations created by this offline run."""
    from playwright.sync_api import expect

    rows = page.locator('.conversation-panel .el-menu-item')
    original = rows.filter(has_text='角色停用后能否登录？')
    question_box = page.get_by_placeholder('针对当前项目资料提问…')
    question_box.fill('删除其他会话不应影响当前草稿')
    deletes = []
    page.on('request', lambda request: deletes.append(request.url) if request.method == 'DELETE' else None)
    dialog = page.locator('.el-message-box').filter(has_text='删除会话')

    original.get_by_role('button', name=re.compile(r'^删除会话')).click()
    expect(dialog).to_contain_text(re.compile(r'(无法|不可)恢复'))
    dialog.get_by_role('button', name='取消', exact=True).click()
    expect(dialog).to_be_hidden()
    expect(rows).to_have_count(3)
    expect(page.locator('.chat-message')).to_have_count(4)
    assert not deletes, 'Cancelling the confirmation must not send a DELETE'

    # Delete B while A is selected: click.stop must preserve both A's history
    # and its unsent draft, and DELETE must not be a mere local list removal.
    rows.nth(1).get_by_role('button', name=re.compile(r'^删除会话')).click()
    dialog.get_by_role('button', name='删除', exact=True).click()
    expect(rows).to_have_count(2)
    expect(original).to_have_class(re.compile(r'\bis-active\b'))
    expect(question_box).to_have_value('删除其他会话不应影响当前草稿')
    expect(page.locator('.chat-message')).to_have_count(4)

    original.get_by_role('button', name=re.compile(r'^删除会话')).click()
    dialog.get_by_role('button', name='删除', exact=True).click()
    expect(rows).to_have_count(1)
    expect(rows.first).to_have_class(re.compile(r'\bis-active\b'))
    expect(question_box).to_have_value('')
    expect(page.locator('.chat-message')).to_have_count(0)
    rows.first.get_by_role('button', name=re.compile(r'^删除会话')).click()
    dialog.get_by_role('button', name='删除', exact=True).click()
    expect(rows).to_have_count(0)
    expect(question_box).to_have_count(0)
    expect(page.locator('.conversation-panel')).to_contain_text('暂无会话')
    assert len(deletes) == 3

    page.reload()
    page.get_by_role('tab', name='知识问答').click()
    expect(rows).to_have_count(0)
    expect(page.locator('.conversation-panel')).to_contain_text('暂无会话')
    page.get_by_role('button', name='新建', exact=True).click()
    expect(rows).to_have_count(1)
    expect(question_box).to_have_value('')
    expect(page.locator('.chat-message')).to_have_count(0)
    expect(page.locator('.chat-panel .task-card')).to_have_count(0)
    expect(page.locator('.chat-messages .el-loading-mask:visible')).to_have_count(0)
    page.screenshot(path=str(artifacts_dir / 'qa-deletion-new-conversation.png'), animations='disabled')


def _assert_database(project):
    from project_knowledge.models import KnowledgeChunk, KnowledgeConversation, KnowledgeDocument, KnowledgeMessage, KnowledgeTask, ManualTestCase

    document = KnowledgeDocument.objects.get(project=project)
    assert document.current_revision.parse_status == "ready"
    assert document.current_revision.index_status == "ready"
    chunk = KnowledgeChunk.objects.get(revision=document.current_revision)
    case = ManualTestCase.objects.get(project=project)
    assert case.title == "停用角色不可登录（已编辑）"
    assert case.sources and case.sources[0]["id"] == str(chunk.id)
    assert KnowledgeConversation.objects.filter(project=project).count() == 1
    assert not KnowledgeMessage.objects.filter(conversation__project=project).exists()
    assert not KnowledgeTask.objects.filter(project=project, kind='answer').exists()
    assert KnowledgeTask.objects.filter(project=project, kind='generate').exists()


def main() -> int:
    parser = argparse.ArgumentParser(description="项目知识库离线真实浏览器联调")
    parser.add_argument(
        "--artifacts-dir",
        type=Path,
        default=BACKEND_DIR / "logs" / "knowledge-browser-check",
        help="保留成功截图和失败截图的 gitignored 目录",
    )
    args = parser.parse_args()
    artifacts_dir = args.artifacts_dir.resolve()
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    if not FRONTEND_DIST.joinpath("index.html").is_file():
        raise SystemExit(f"前端构建产物不存在: {FRONTEND_DIST}")
    if not CHROME.is_file() or not os.access(CHROME, os.X_OK):
        raise SystemExit(f"未找到可执行 Chrome: {CHROME}；可用 AITS_TEST_CHROME_EXECUTABLE 覆盖")

    original_connect, original_connect_ex = socket.socket.connect, socket.socket.connect_ex
    with tempfile.TemporaryDirectory(prefix="aits-project-knowledge-browser-") as temporary, patch.object(socket.socket, "connect", _loopback_only(original_connect)), patch.object(socket.socket, "connect_ex", _loopback_only(original_connect_ex)):
        root = Path(temporary)
        user, project, model, token = _bootstrap_database(root)
        from aits_backend.wsgi import application as django_app
        from project_knowledge import indexing, llm, tasks

        server = make_server("127.0.0.1", 0, _static_or_django(django_app), handler_class=_QuietHandler)
        origin = f"http://127.0.0.1:{server.server_port}"
        server_thread = threading.Thread(target=server.serve_forever, daemon=True)
        with patch.object(tasks.execute_knowledge_task, "delay", side_effect=_inline_delay), patch.object(indexing, "index_revision", side_effect=lambda revision: revision.chunks.count()), patch.object(indexing, "search_chunks", side_effect=_fake_search), patch.object(llm, "stream_call", side_effect=_fake_stream_call):
            server_thread.start()
            try:
                _ui_flow(origin, root, artifacts_dir, token, user, project, model)
                _assert_database(project)
            finally:
                server.shutdown()
                server.server_close()
                server_thread.join(timeout=5)
    print(f"project knowledge browser integration: PASS (artifacts: {artifacts_dir})")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        traceback.print_exc()
        raise SystemExit(1)
