"""Readable, text-only Excel exports for reviewed or unreviewed manual cases."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_workbook(cases):
    book = Workbook()
    sheet = book.active
    sheet.title = '手工测试用例'
    sheet.append(['标题', '模块', '类型', '前置条件', '测试数据', '测试步骤', '预期结果', '来源', '测试点', '待确认项', '审核状态'])
    for case in cases:
        sources = []
        for source in case.get('sources') or []:
            version = (source.get('version') or {}).get('number', source.get('revision_id', '未知'))
            state = '当前可用' if source.get('active') else '来源已更新、停用或删除'
            sources.append(f'{source.get("file_name", "来源")} · {source.get("heading", "")} · 版本 {version}（{state}）\n{source.get("content", "")}')
        steps = case.get('steps') or []
        values = [case['title'], case.get('module', '未分类'), case.get('test_type', ''),
                  '\n'.join(case.get('preconditions') or []), case.get('test_data', ''),
                  '\n'.join(f'{i}. {step.get("action", "")}' for i, step in enumerate(steps, 1)),
                  '\n'.join(f'{i}. {step.get("expected", "")}' for i, step in enumerate(steps, 1)),
                  '\n\n'.join(sources), '\n'.join(case.get('test_point_ids') or []),
                  '\n'.join(case.get('pending_questions') or []),
                  '已审核' if case.get('review_status') == 'reviewed' else '未审核']
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('', str(value)) for value in values])
        # Explicit string cells prevent formula interpretation without changing
        # the user's actual text (including leading spaces, =, +, - and @).
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
        cell.alignment = Alignment(vertical='center')
    for index, width in enumerate([36, 16, 12, 30, 28, 48, 48, 55, 25, 36, 12], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
