from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
import uuid
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase

from projects.models import Project, UploadedFile
from project_knowledge import indexing
from project_knowledge.indexing import (
    ProjectKnowledgeIndex,
    ProjectKnowledgeSearchUnavailable,
    clear_project_knowledge_index_cache,
)
from project_knowledge.models import DocumentRevision, KnowledgeChunk, KnowledgeDocument
from project_knowledge.parsing import parse_document


class DocumentParsingTests(SimpleTestCase):
    def test_markdown_sections_are_source_addressable_and_long_content_is_preserved(self):
        content = "# 登录规则\n" + ("登录失败后锁定。" * 12) + "\n## 例外\n管理员不受此限制。\n"
        with TemporaryDirectory() as directory, patch("project_knowledge.parsing.MAX_CHUNK_CHARS", 24):
            source = Path(directory) / "rules.md"
            source.write_text(content, encoding="utf-8")
            chunks = parse_document(str(source), "rules.md")

        login_chunks = [chunk for chunk in chunks if chunk["heading"] == "登录规则"]
        self.assertGreater(len(login_chunks), 1)
        self.assertTrue(all(chunk["content"].startswith("登录规则\n") for chunk in login_chunks))
        self.assertEqual(
            "".join(chunk["content"].removeprefix("登录规则\n") for chunk in login_chunks),
            content[len("# 登录规则"):content.index("## 例外")],
        )
        self.assertEqual([chunk["ordinal"] for chunk in chunks], list(range(len(chunks))))
        self.assertTrue(all(chunk["location"]["character_end"] > chunk["location"]["character_start"] for chunk in chunks))

    def test_docx_and_xlsx_include_real_paragraph_or_row_locations(self):
        from docx import Document

        with TemporaryDirectory() as directory:
            directory_path = Path(directory)
            docx_path = directory_path / "rules.docx"
            document = Document()
            document.add_heading("账户规则", level=1)
            document.add_paragraph("密码重置后，旧会话失效。")
            document.save(docx_path)

            docx_chunks = parse_document(str(docx_path), "rules.docx")
            self.assertEqual(docx_chunks[0]["heading"], "账户规则")
            self.assertEqual(docx_chunks[0]["content"], "账户规则\n密码重置后，旧会话失效。")
            self.assertEqual(docx_chunks[0]["location"]["paragraph_start"], 2)

            xlsx_path = directory_path / "rules.xlsx"
            _write_minimal_xlsx(xlsx_path)
            xlsx_chunks = parse_document(str(xlsx_path), "rules.xlsx")

        self.assertEqual([chunk["content"] for chunk in xlsx_chunks], ["字段 | 规则", "角色 | 停用后禁止登录"])
        self.assertEqual(xlsx_chunks[1]["location"], {
            "kind": "xlsx", "sheet": "权限", "row_start": 2, "row_end": 2,
            "character_start": 0, "character_end": len("角色 | 停用后禁止登录"),
        })

    def test_text_pdf_has_page_location_and_scanned_pdf_is_rejected(self):
        pages = [SimpleNamespace(extract_text=lambda: ""), SimpleNamespace(extract_text=lambda: "第1章 权限\n只读成员不可编辑")]
        with TemporaryDirectory() as directory:
            source = Path(directory) / "rules.pdf"
            source.write_bytes(b"placeholder")
            with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
                chunks = parse_document(str(source), "rules.pdf")
            self.assertEqual(chunks[0]["location"]["page_start"], 2)
            self.assertEqual(chunks[0]["heading"], "第1章 权限")

            with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "")])):
                with self.assertRaisesRegex(ValueError, "扫描 PDF"):
                    parse_document(str(source), "rules.pdf")

    def test_legacy_doc_is_explicitly_rejected(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "legacy.doc"
            source.write_bytes(b"old binary word")
            with self.assertRaisesRegex(ValueError, "不支持的文档格式"):
                parse_document(str(source), "legacy.doc")

    def test_numbered_text_rules_and_heading_only_rule_remain_semantic_content(self):
        with TemporaryDirectory() as directory:
            directory_path = Path(directory)
            numbered = directory_path / "rules.txt"
            numbered.write_text("1. 用户名不能重复\n2. 密码不少于8位", encoding="utf-8")
            numbered_chunks = parse_document(str(numbered), "rules.txt")

            heading_only = directory_path / "heading-only.md"
            heading_only.write_text("# 管理员不得绕过审批", encoding="utf-8")
            heading_chunks = parse_document(str(heading_only), "heading-only.md")

        self.assertEqual([chunk["content"] for chunk in numbered_chunks], ["1. 用户名不能重复", "2. 密码不少于8位"])
        self.assertEqual(heading_chunks[0]["content"], "管理员不得绕过审批")

    def test_docx_table_rows_keep_body_order_and_table_row_locations(self):
        from docx import Document

        with TemporaryDirectory() as directory:
            source = Path(directory) / "table-rules.docx"
            document = Document()
            document.add_heading("权限规则", level=1)
            document.add_paragraph("下列规则适用于成员。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "角色"
            table.cell(0, 1).text = "权限"
            table.cell(1, 0).text = "只读成员"
            table.cell(1, 1).text = "不可编辑"
            document.add_paragraph("变更立即生效。")
            document.save(source)
            chunks = parse_document(str(source), "table-rules.docx")

        contents = [chunk["content"] for chunk in chunks]
        table_rule_index = contents.index("权限规则\n只读成员 | 不可编辑")
        self.assertLess(contents.index("权限规则\n下列规则适用于成员。"), table_rule_index)
        self.assertLess(table_rule_index, contents.index("权限规则\n变更立即生效。"))
        self.assertEqual(chunks[table_rule_index]["location"]["table_index"], 1)
        self.assertEqual(chunks[table_rule_index]["location"]["row_start"], 2)


class _KnowledgeFixtureMixin:
    def _user(self, username):
        from django.contrib.auth import get_user_model

        return get_user_model().objects.create_user(
            username=username,
            email=f"{username}@example.test",
            password="not-a-production-password",
        )

    def _revision_with_chunk(self, project, heading, content, *, make_current=True, number=1):
        uploaded = UploadedFile.objects.create(
            original_name=f"{heading}.txt",
            file=SimpleUploadedFile(f"{heading}.txt", content.encode("utf-8")),
            file_size=len(content.encode("utf-8")),
            file_hash=uuid.uuid4().hex,
            file_type="txt",
            upload_status="uploaded",
            uploaded_by=self.user,
            project=project,
        )
        document = KnowledgeDocument.objects.create(
            project=project,
            name=f"{heading}.txt",
            created_by=self.user,
        )
        revision = DocumentRevision.objects.create(
            document=document,
            number=number,
            uploaded_file=uploaded,
            parse_status="ready",
            index_status="ready",
        )
        if make_current:
            document.current_revision = revision
            document.save(update_fields=["current_revision"])
        chunk = KnowledgeChunk.objects.create(
            revision=revision,
            ordinal=0,
            section_key="section-0001",
            heading=heading,
            content=content,
            location={"kind": "text", "character_start": 0, "character_end": len(content)},
        )
        return revision, chunk


class ProjectKnowledgeIndexTests(_KnowledgeFixtureMixin, TestCase):
    def setUp(self):
        self.user = self._user("owner")
        self.project = Project.objects.create(name="项目一", project_type="web", created_by=self.user)
        self.other_project = Project.objects.create(name="项目二", project_type="web", created_by=self.user)
        self.vector = _FakeVectorStore()
        self.index = ProjectKnowledgeIndex(
            config=SimpleNamespace(vector_db_type="chroma"),
            embeddings=object(),
            vector_store=self.vector,
        )

    def test_index_retry_uses_chunk_uuid_upserts_and_delete_is_revision_scoped(self):
        revision, chunk = self._revision_with_chunk(self.project, "登录规则", "角色停用后禁止登录")
        self.index.index_revision(revision)
        self.index.index_revision(revision)

        self.assertEqual(set(self.vector.documents), {str(chunk.id)})
        self.assertEqual(self.vector.add_calls, 2)
        self.assertEqual(self.vector.documents[str(chunk.id)].metadata["revision_id"], str(revision.id))
        self.assertEqual(self.vector.documents[str(chunk.id)].page_content, "登录规则\n角色停用后禁止登录")

        with patch("project_knowledge.indexing.get_project_knowledge_index", return_value=self.index):
            indexing.delete_revision_index(str(revision.id))
        self.assertEqual(self.vector.documents, {})

    def test_milvus_retry_uses_its_explicit_upsert(self):
        revision, chunk = self._revision_with_chunk(self.project, "Milvus 规则", "同一分片不可重复")
        milvus = _FakeMilvusVectorStore()
        index = ProjectKnowledgeIndex(
            config=SimpleNamespace(vector_db_type="milvus"),
            embeddings=object(),
            vector_store=milvus,
        )

        index.index_revision(revision)
        index.index_revision(revision)

        self.assertEqual(set(milvus.documents), {str(chunk.id)})
        self.assertEqual(milvus.upsert_calls, 2)

    def test_search_enforces_project_current_version_and_database_post_filter(self):
        current_revision, current_chunk = self._revision_with_chunk(self.project, "登录规则", "角色停用后禁止登录")
        old_revision, _ = self._revision_with_chunk(self.project, "旧规则", "旧版本不得返回", make_current=False, number=2)
        other_revision, _ = self._revision_with_chunk(self.other_project, "登录规则", "另一个项目不得返回")
        for revision in (current_revision, old_revision, other_revision):
            self.index.index_revision(revision)

        with patch("project_knowledge.indexing.get_project_knowledge_index", return_value=self.index):
            result = indexing.search_chunks(
                self.project.id,
                [str(current_revision.id), str(old_revision.id), str(other_revision.id)],
                "登录规则",
                top_k=8,
            )

        self.assertEqual([item["id"] for item in result], [str(current_chunk.id)])
        self.assertEqual(result[0]["file_name"], "登录规则.txt")
        self.assertEqual(result[0]["document_id"], current_revision.document_id)
        self.assertIn("$and", self.vector.last_filter)
        self.assertNotIn("旧版本", result[0]["content"])

    def test_vector_filter_error_fails_closed(self):
        revision, _ = self._revision_with_chunk(self.project, "登录规则", "角色停用后禁止登录")
        self.index.index_revision(revision)
        self.vector.search_error = RuntimeError("backend filter rejected")

        with patch("project_knowledge.indexing.get_project_knowledge_index", return_value=self.index):
            with self.assertRaisesRegex(ProjectKnowledgeSearchUnavailable, "项目资料检索暂时不可用"):
                indexing.search_chunks(self.project.id, [str(revision.id)], "登录", top_k=8)

        self.assertIsNotNone(self.vector.last_filter)

class ProjectKnowledgeIndexCacheTests(SimpleTestCase):
    def setUp(self):
        clear_project_knowledge_index_cache()

    def tearDown(self):
        clear_project_knowledge_index_cache()

    def test_cache_reuses_only_the_same_configuration_fingerprint(self):
        first_config = _cache_config(updated_at="2026-09-06T10:00:00+08:00")
        changed_config = _cache_config(updated_at="2026-09-06T10:01:00+08:00")
        first_index, changed_index, cleared_index = object(), object(), object()

        with patch("project_knowledge.indexing._load_active_rag_config", side_effect=[first_config, first_config, changed_config, changed_config]), patch(
            "project_knowledge.indexing.ProjectKnowledgeIndex",
            side_effect=[first_index, changed_index, cleared_index],
        ) as constructor:
            self.assertIs(indexing.get_project_knowledge_index(), first_index)
            self.assertIs(indexing.get_project_knowledge_index(), first_index)
            self.assertIs(indexing.get_project_knowledge_index(), changed_index)
            clear_project_knowledge_index_cache()
            self.assertIs(indexing.get_project_knowledge_index(), cleared_index)

        self.assertEqual(constructor.call_count, 3)


class ChromaIndexIntegrationTests(_KnowledgeFixtureMixin, TestCase):
    def setUp(self):
        self.user = self._user("chroma-owner")
        self.project = Project.objects.create(name="Chroma 项目一", project_type="web", created_by=self.user)
        self.other_project = Project.objects.create(name="Chroma 项目二", project_type="web", created_by=self.user)

    def test_local_chroma_filters_upserts_and_deletes_without_real_embeddings(self):
        revision, chunk = self._revision_with_chunk(self.project, "登录规则", "角色停用后禁止登录")
        other_revision, _ = self._revision_with_chunk(self.other_project, "登录规则", "另一项目内容")
        with TemporaryDirectory() as directory:
            actual_index = ProjectKnowledgeIndex(
                config=SimpleNamespace(vector_db_type="chroma", chroma_persist_directory=directory),
                embeddings=_FixedEmbeddings(),
            )
            actual_index.index_revision(revision)
            actual_index.index_revision(revision)
            actual_index.index_revision(other_revision)

            self.assertEqual(actual_index.vector_store._collection.count(), 2)
            result = actual_index.search(self.project.id, [str(revision.id)], "登录规则", top_k=8)
            self.assertEqual([item["id"] for item in result], [str(chunk.id)])

            actual_index.delete_revision(str(revision.id))
            self.assertEqual(actual_index.vector_store._collection.count(), 1)
            self.assertEqual(actual_index.search(self.project.id, [str(revision.id)], "登录规则", top_k=8), [])


class _FakeCollection:
    def __init__(self, store):
        self.store = store

    def get(self, *, where, include):
        return {
            "ids": [
                vector_id for vector_id, document in self.store.documents.items()
                if document.metadata.get("revision_id") == where.get("revision_id")
            ]
        }

    def delete(self, *, ids):
        for vector_id in ids:
            self.store.documents.pop(vector_id, None)


class _FakeVectorStore:
    def __init__(self):
        self.documents = {}
        self.add_calls = 0
        self.last_filter = None
        self.search_error = None
        self._collection = _FakeCollection(self)

    def add_documents(self, documents, *, ids):
        self.add_calls += 1
        self.documents.update(dict(zip(ids, documents)))

    def similarity_search_with_score(self, query, *, k, filter):
        self.last_filter = filter
        if self.search_error:
            raise self.search_error
        # The fake intentionally ignores the filter.  The production code must
        # still reject all non-current/non-project records with its DB check.
        return [(document, 0.1) for document in self.documents.values()][:k]


class _FakeMilvusVectorStore:
    def __init__(self):
        self.documents = {}
        self.upsert_calls = 0

    def upsert(self, *, ids, documents):
        self.upsert_calls += 1
        self.documents.update(dict(zip(ids, documents)))


class _FixedEmbeddings:
    """Stable local vectors for the real Chroma test; no model files are read."""

    @staticmethod
    def _vector(text):
        return [float(len(text) % 7), float(text.count("登录")), float(text.count("规则")), 1.0]

    def embed_documents(self, texts):
        return [self._vector(text) for text in texts]

    def embed_query(self, text):
        return self._vector(text)


def _cache_config(*, updated_at):
    return SimpleNamespace(
        id=7,
        updated_at=updated_at,
        vector_db_type="chroma",
        embedding_model="BAAI/test-model",
        chroma_persist_directory="/tmp/project-knowledge-cache-test",
        milvus_host="localhost",
        milvus_port=19530,
        milvus_collection_name="knowledge_base",
    )


def _write_minimal_xlsx(path: Path):
    with zipfile.ZipFile(path, "w") as workbook:
        workbook.writestr("xl/workbook.xml", """<?xml version='1.0'?>
            <workbook xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
             xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>
              <sheets><sheet name='权限' sheetId='1' r:id='rId1'/></sheets>
            </workbook>""")
        workbook.writestr("xl/_rels/workbook.xml.rels", """<?xml version='1.0'?>
            <Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>
              <Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet' Target='worksheets/sheet1.xml'/>
            </Relationships>""")
        workbook.writestr("xl/sharedStrings.xml", """<?xml version='1.0'?>
            <sst xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main' count='4' uniqueCount='4'>
              <si><t>字段</t></si><si><t>规则</t></si><si><t>角色</t></si><si><t>停用后禁止登录</t></si>
            </sst>""")
        workbook.writestr("xl/worksheets/sheet1.xml", """<?xml version='1.0'?>
            <worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData>
              <row r='1'><c r='A1' t='s'><v>0</v></c><c r='B1' t='s'><v>1</v></c></row>
              <row r='2'><c r='A2' t='s'><v>2</v></c><c r='B2' t='s'><v>3</v></c></row>
            </sheetData></worksheet>""")
