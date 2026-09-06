from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase

from project_knowledge.parsing import parse_document


class ParsingBoundaryTests(SimpleTestCase):
    def test_nested_markdown_titles_retain_module_scope(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / 'rules.md'
            source.write_text('# 用户\n## 名称\n不得重复。\n# 角色\n## 名称\n允许重复。', encoding='utf-8')
            chunks = parse_document(str(source), source.name)
        self.assertIn('用户 > 名称\n\n不得重复。\n', [item['content'] for item in chunks])
        self.assertIn('角色 > 名称\n\n允许重复。', [item['content'] for item in chunks])

    def test_nested_word_titles_retain_scope_and_table_rules(self):
        from docx import Document
        with TemporaryDirectory() as folder:
            source = Path(folder) / 'rules.docx'
            document = Document()
            document.add_heading('用户', level=1)
            document.add_heading('名称', level=2)
            document.add_paragraph('不得重复。')
            document.save(source)
            chunks = parse_document(str(source), source.name)
        self.assertIn('用户 > 名称\n不得重复。', [item['content'] for item in chunks])

    def test_mixed_scanned_pdf_does_not_silently_claim_complete(self):
        pages = [SimpleNamespace(extract_text=lambda: '用户不能重复', images=[]),
                 SimpleNamespace(extract_text=lambda: '', images=['scanned-image'])]
        with TemporaryDirectory() as folder:
            source = Path(folder) / 'mixed.pdf'
            source.write_bytes(b'offline-pdf-fixture')
            with patch('pypdf.PdfReader', return_value=SimpleNamespace(pages=pages)):
                with self.assertRaisesRegex(ValueError, '第 2 页.*正文不完整'):
                    parse_document(str(source), source.name)

    def test_oversized_office_archive_is_rejected_before_xml_parsing(self):
        import zipfile
        with TemporaryDirectory() as folder:
            source = Path(folder) / 'rules.xlsx'
            with zipfile.ZipFile(source, 'w', zipfile.ZIP_DEFLATED) as package:
                package.writestr('xl/workbook.xml', 'x' * 100)
            with patch('project_knowledge.parsing.MAX_OFFICE_EXPANDED_BYTES', 50):
                with self.assertRaisesRegex(ValueError, '展开后体积过大'):
                    parse_document(str(source), source.name)
